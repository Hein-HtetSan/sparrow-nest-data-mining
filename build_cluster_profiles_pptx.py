from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "presentation"
OUT.mkdir(exist_ok=True)
GREEN = RGBColor(31, 74, 51)
LIGHT = RGBColor(238, 244, 239)
WARM = RGBColor(154, 97, 62)
INK = RGBColor(25, 30, 26)
MUTED = RGBColor(91, 99, 93)
WHITE = RGBColor(255, 255, 255)

profiles = pd.read_csv(ROOT / "outputs/cluster_profiles.csv").set_index("Cluster")
nests = pd.read_csv(ROOT / "outputs/cluster_nest_profiles.csv").set_index("Cluster")
names = {
    0: "Moderate-altitude, noisy",
    1: "Warm, wet, low-altitude",
    2: "High-altitude, dry, wide",
    3: "Highest-altitude, cool, low-wind",
    4: "Very tall buildings",
}
descriptions = {
    0: "Lower buildings, moderate altitude, highest noise and above-average precipitation",
    1: "Lowest altitude, warmest and wettest conditions, relatively tall and narrow buildings",
    2: "Lowest and widest buildings, high altitude, driest conditions and slightly higher wind",
    3: "Highest altitude, coolest temperature and lowest wind, relatively low buildings",
    4: "Tallest buildings by a large margin, otherwise moderate environmental conditions",
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def background(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def textbox(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def title(slide, text, number):
    textbox(slide, f"CLUSTER ANALYSIS  /  {number:02d}", 0.7, 0.35, 4.0, 0.3, 11, GREEN, True)
    textbox(slide, text, 0.7, 0.78, 11.9, 0.75, 30, INK, True)


def table(slide, values, x, y, w, h, widths=None, font_size=14):
    rows, cols = len(values), len(values[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if widths:
        for i, width in enumerate(widths):
            tbl.columns[i].width = Inches(width)
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = GREEN if r == 0 else (LIGHT if r % 2 else WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(font_size if r else font_size - 1)
                paragraph.font.bold = r == 0 or c == 0
                paragraph.font.color.rgb = WHITE if r == 0 else INK
    return shape


slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, GREEN)
textbox(slide, "DATA MINING PROJECT", 0.8, 0.65, 4.0, 0.35, 13, RGBColor(190, 218, 198), True)
textbox(slide, "Cluster Characteristic Profiles", 0.8, 1.55, 10.8, 1.4, 42, WHITE, True)
textbox(slide, "Eurasian Tree Sparrow building and environmental groups", 0.82, 3.1, 9.5, 0.7, 22, RGBColor(222, 234, 225))
textbox(slide, "501 training records\nK-means k = 5\nSeven standardized features", 0.82, 5.2, 4.2, 1.1, 16, RGBColor(190, 218, 198), True)

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide)
title(slide, "How the profiles were formed", 2)
slide.shapes.add_picture(str(ROOT / "outputs/cluster_pca.png"), Inches(0.75), Inches(1.7), width=Inches(7.3))
textbox(slide, "K-means input", 8.45, 1.78, 3.5, 0.4, 18, GREEN, True)
textbox(slide, "Building height and width\nAltitude and average noise\nTemperature, precipitation and wind", 8.45, 2.3, 4.1, 1.45, 17, INK)
textbox(slide, "Why standardize?", 8.45, 4.0, 3.5, 0.4, 18, GREEN, True)
textbox(slide, "StandardScaler gives every feature equal influence despite different units.", 8.45, 4.52, 4.0, 0.9, 17, INK)
textbox(slide, "NestCount was excluded from K-means and compared after clustering.", 8.45, 5.72, 4.0, 0.8, 16, WARM, True)

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide)
title(slide, "Five building and environmental profiles", 3)
overview = [["Cluster profile", "Records", "Share", "Defining characteristics", "Mean nests"]]
for cluster in range(5):
    overview.append([
        f"{cluster}  {names[cluster]}",
        int(profiles.at[cluster, "Records"]),
        f"{profiles.at[cluster, 'Records'] / 501:.1%}",
        descriptions[cluster],
        f"{nests.at[cluster, 'MeanNestCount']:.2f}",
    ])
table(slide, overview, 0.6, 1.65, 12.1, 4.8, [2.45, 0.8, 0.75, 6.9, 1.2], 13)
textbox(slide, "Cluster 1 has the highest observed nest count. Cluster 3 has the lowest.", 0.72, 6.62, 11.8, 0.35, 15, GREEN, True)

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide)
title(slide, "Original-scale cluster means", 4)
means = [["Cluster", "Height", "Width", "Altitude", "Noise", "Temp", "Precip.", "Wind"]]
for cluster in range(5):
    means.append([
        f"{cluster}  {names[cluster]}",
        f"{profiles.at[cluster, 'BuildingHeight']:.2f}",
        f"{profiles.at[cluster, 'BuildingWidth']:.2f}",
        f"{profiles.at[cluster, 'Altitude']:.2f}",
        f"{profiles.at[cluster, 'AveNoise']:.2f}",
        f"{profiles.at[cluster, 'temp']:.2f}",
        f"{profiles.at[cluster, 'prec']:.2f}",
        f"{profiles.at[cluster, 'win']:.2f}",
    ])
table(slide, means, 0.55, 1.75, 12.2, 4.55, [3.1, 1.05, 1.05, 1.25, 1.0, 1.0, 1.15, 1.0], 12)
textbox(slide, "These means describe the center of each cluster in the variables' original units.", 0.7, 6.55, 11.7, 0.45, 15, MUTED)

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide)
title(slide, "Interpretation and limits", 5)
textbox(slide, "Strongest contrasts", 0.8, 1.75, 4.0, 0.5, 22, GREEN, True)
textbox(slide, "Cluster 1 combines very low altitude with the warmest and wettest conditions.\n\nCluster 3 combines the highest altitude with the coolest temperature and lowest wind.\n\nCluster 4 separates mainly through very tall buildings.", 0.8, 2.45, 5.5, 3.4, 19, INK)
textbox(slide, "What the result supports", 7.0, 1.75, 4.3, 0.5, 22, GREEN, True)
textbox(slide, "The profiles summarize similarity within the 501 training records. Nest differences are descriptive because NestCount did not form the clusters.\n\nSilhouette = 0.267 indicates overlap between groups. The clusters should not be presented as sharply separated natural classes or causal effects.", 7.0, 2.45, 5.2, 3.4, 19, INK)
textbox(slide, "K-means provides interpretable environmental profiles, but the modest silhouette requires cautious interpretation.", 0.8, 6.25, 11.7, 0.5, 16, WARM, True)
textbox(slide, "Live model lab: https://sparrow-nest-data-mining.streamlit.app", 0.8, 6.72, 11.7, 0.35, 13, GREEN, True)

prs.save(OUT / "Sparrow_Cluster_Profiles.pptx")
